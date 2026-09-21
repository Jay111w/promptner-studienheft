"""Erzeugt die strukturierte Word-Korrekturuebersicht (Tabelle).

Spalten: ID, PDF-Seite, Heft-Seite, Absatz, Fehlerhafter Ausdruck,
Korrektur, Kategorie, Kommentar/Erklaerung. Die ID (z. B. K-001) stellt die
konsistente Verbindung zur PDF-Annotation her.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

from korrektor.config import get_logger
from korrektor.domain import Category, Correction
from korrektor.errors import ErrorCode, ExportError

log = get_logger("services.export.word_exporter")

_HEADERS = [
    "ID",
    "PDF-Seite",
    "Heft-Seite",
    "Absatz",
    "Fehlerhafter Ausdruck",
    "Korrektur",
    "Kat.",
    "Kommentar / Erklaerung",
]


def export_word(
    corrections: list[Correction],
    output_path: str | Path,
    title: str = "Korrekturuebersicht",
) -> Path:
    """Schreibt die Korrekturen als Word-Tabelle.

    Raises:
        ExportError: Bei Schreibfehlern.
    """
    output_path = Path(output_path)
    try:
        document = Document()
        heading = document.add_heading(title, level=1)
        heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

        _add_legend(document)

        table = document.add_table(rows=1, cols=len(_HEADERS))
        table.style = "Light Grid Accent 1"
        for cell, header in zip(table.rows[0].cells, _HEADERS, strict=True):
            run = cell.paragraphs[0].add_run(header)
            run.bold = True
            run.font.size = Pt(9)

        for c in corrections:
            _add_row(table, c)

        output_path.parent.mkdir(parents=True, exist_ok=True)
        document.save(output_path)
    except Exception as exc:  # noqa: BLE001
        raise ExportError(
            ErrorCode.WORD_EXPORT_FAILED,
            context={"output": str(output_path)},
            cause=exc,
        ) from exc

    log.info(
        "Word-Uebersicht geschrieben: %s (%d Zeilen).",
        output_path.name,
        len(corrections),
    )
    return output_path


def _add_legend(document) -> None:
    """Schreibt die farbige Kategorie-Legende ueber die Tabelle."""
    paragraph = document.add_paragraph()
    intro = paragraph.add_run("Legende:   ")
    intro.bold = True
    for category in (Category.A, Category.B, Category.C):
        marker = paragraph.add_run("■ ")  # gefuelltes Quadrat
        marker.bold = True
        marker.font.color.rgb = RGBColor(*category.rgb)
        label = paragraph.add_run(f"{category.value} = {category.label}    ")
        label.font.size = Pt(9)


def _add_row(table, correction: Correction) -> None:
    printed = str(correction.printed_page) if correction.printed_page is not None else "-"
    values = [
        correction.id,
        str(correction.pdf_page + 1),
        printed,
        correction.paragraph or "-",
        correction.original_text,
        correction.corrected_text or "-",
        correction.category.value,
        _comment(correction),
    ]
    cells = table.add_row().cells
    category_col = 6
    for col, (cell, value) in enumerate(zip(cells, values, strict=True)):
        if col == category_col:
            # Kategorie-Buchstabe farbig hervorheben (Gelb=A, Blau=B, Gruen=C).
            run = cell.paragraphs[0].add_run(value)
            run.bold = True
            run.font.color.rgb = RGBColor(*correction.category.rgb)
        else:
            cell.text = value


def _comment(correction: Correction) -> str:
    parts: list[str] = [f"Aktion: {correction.action_type.value}"]
    if correction.description:
        parts.append(f"Fehler: {correction.description}")
    if correction.reason:
        parts.append(f"Begruendung: {correction.reason}")
    if correction.editorial_note:
        parts.append(f"Hinweis: {correction.editorial_note}")
    return "\n".join(parts)
