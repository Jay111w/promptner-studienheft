"""Integrationstests fuer die Exporter (PDF-Annotation, Word, Changelog)."""

import fitz
import pytest
from docx import Document
from tests.fixtures.pdf_factory import make_pdf

from korrektor.domain import (
    ActionType,
    Category,
    Correction,
    TextLocation,
)
from korrektor.services.export import (
    export_annotated_pdf,
    export_changelog,
    export_word,
)
from korrektor.services.pdf.extractor import PdfDocument

PAGE_TEXT = "Hier steht das Wort Veranstalltung mit einem Fehler."


@pytest.fixture
def source_pdf(tmp_path):
    return make_pdf(tmp_path / "quelle.pdf", [PAGE_TEXT])


@pytest.fixture
def located_correction(source_pdf):
    """Eine Korrektur mit echten Trefferkoordinaten aus dem PDF."""
    with PdfDocument(source_pdf) as doc:
        rects = doc.search_rects(0, "Veranstalltung")
    return Correction(
        id="K-001",
        pdf_page=0,
        printed_page=1,
        paragraph="Absatz 1",
        original_text="Veranstalltung",
        corrected_text="Veranstaltung",
        category=Category.A,
        action_type=ActionType.REPLACE,
        description="Doppel-l",
        reason="Rechtschreibung (Duden)",
        location=TextLocation(pdf_page=0, quote="Veranstalltung", rects=rects),
    )


@pytest.mark.integration
def test_pdf_export_writes_highlight_with_comment(source_pdf, located_correction, tmp_path):
    out = export_annotated_pdf(source_pdf, [located_correction], tmp_path / "out.pdf")
    assert out.is_file()

    doc = fitz.open(out)
    try:
        annots = list(doc[0].annots())
        assert len(annots) >= 1
        contents = " ".join(a.info.get("content", "") for a in annots)
        assert "[K-001]" in contents
        assert "Kategorie A" in contents
    finally:
        doc.close()


@pytest.mark.integration
def test_pdf_export_fallback_note_when_no_rects(source_pdf, tmp_path):
    correction = Correction(
        id="K-002",
        pdf_page=0,
        original_text="nicht auffindbar",
        corrected_text="x",
        category=Category.B,
        action_type=ActionType.REPLACE,
        location=TextLocation(pdf_page=0, quote="nicht auffindbar", rects=[]),
    )
    out = export_annotated_pdf(source_pdf, [correction], tmp_path / "out.pdf")
    doc = fitz.open(out)
    try:
        assert len(list(doc[0].annots())) >= 1
    finally:
        doc.close()


@pytest.mark.integration
def test_word_export_has_header_and_row(located_correction, tmp_path):
    out = export_word([located_correction], tmp_path / "uebersicht.docx")
    assert out.is_file()

    document = Document(out)
    table = document.tables[0]
    assert table.rows[0].cells[0].text == "ID"
    # Kopfzeile + eine Datenzeile
    assert len(table.rows) == 2
    row = table.rows[1]
    assert row.cells[0].text == "K-001"
    assert row.cells[4].text == "Veranstalltung"
    assert row.cells[6].text == "A"


@pytest.mark.integration
def test_changelog_pdf_is_valid(located_correction, tmp_path):
    out = export_changelog([located_correction], tmp_path / "changelog.pdf", "Heft 1")
    assert out.is_file()

    doc = fitz.open(out)
    try:
        assert doc.page_count >= 1
        text = doc[0].get_text("text")
    finally:
        doc.close()
    assert "Aenderungsdokumentation" in text
    assert "K-001" in text


@pytest.mark.integration
def test_word_export_handles_empty_list(tmp_path):
    out = export_word([], tmp_path / "leer.docx")
    document = Document(out)
    assert len(document.tables[0].rows) == 1  # nur Kopfzeile


@pytest.mark.integration
def test_word_export_contains_color_legend(located_correction, tmp_path):
    out = export_word([located_correction], tmp_path / "mit_legende.docx")
    document = Document(out)
    full_text = "\n".join(p.text for p in document.paragraphs)
    assert "Legende" in full_text
    assert "A = Sicherer Fehler" in full_text
    assert "B = Sprachliche Verbesserung" in full_text
    assert "C = Stil-/Redaktionsentscheidung" in full_text
